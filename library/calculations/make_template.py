"""Genera la plantilla de entrada del motor (src/input_motor_AEA_V2.xlsx) a partir del Excel original.

El original no se modifica: se leen sus valores y se escribe un libro nuevo, con las columnas en
orden secuencial, las de instalación, listas desplegables y formato sobrio. Requiere openpyxl
(solo para generar la plantilla; el motor lee el xlsx con su lector propio).
Uso: python3 -B make_template.py ../../src/input_for_sld_V1.xlsx ../../src/input_motor_AEA_V2.xlsx
"""
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

import aea_ampacity as A

FUENTE = "Calibri"
TINTA, OSCURO, MEDIO, CLARO, LINEA, SUAVE = "1F2937", "1F2937", "64748B", "E5E7EB", "D1D5DB", "F8FAFC"
ACENTO = "0F766E"

M = list(A.METODOS_DESCRIPCION.values())
LISTAS = {
    "UNIDAD_CARGA": ["A", "W", "kW", "MW", "VA", "kVA", "MVA", "HP"],
    "FASE": ["1F", "3F"],
    "TOPOLOGY": ["single", "dol", "dol_ol", "star_delta", "softstarter", "vfd", "dol_rcd", "manual"],
    "KIND": ["auto", "mcb", "mccb", "rcbo", "rccb"],
    "MATERIAL": ["Cobre", "Aluminio"],
    "AISLACION": ["PVC", "XLPE"],
    "CLASE": ["2", "5"],
    "METODO": M,
}
# (columna, grupo, obligatoria, ancho, defecto para filas existentes, ayuda)
COLUMNAS = [
    ("TAG", "IDENTIFICACIÓN", True, 10, None, "Identificador único del circuito (ej. C-01)."),
    ("DESCRIPCION", "IDENTIFICACIÓN", True, 34, None, "Descripción del circuito."),
    ("DESTINO", "IDENTIFICACIÓN", True, 26, None, "Carga o tablero que alimenta."),
    ("PARENT_TAG", "IDENTIFICACIÓN", False, 12, None, "TAG del circuito padre; vacío si cuelga de la barra."),
    ("CARGA", "CARGA", True, 9, None, "Valor numérico. Vacío en un padre: se suma la de sus hijos."),
    ("UNIDAD_CARGA", "CARGA", True, 9, None, "A, W, kW, MW, VA, kVA, MVA o HP."),
    ("FASE", "CARGA", False, 7, None, "1F o 3F (defecto 3F)."),
    ("FP", "CARGA", False, 7, None, "Factor de potencia; vacío usa FP_DEFAULT del tablero."),
    ("RENDIMIENTO", "CARGA", False, 13, None, "Rendimiento del motor (0 a 1); se usa si la carga está en HP o kW de eje."),
    ("TOPOLOGY", "CARGA", False, 12, None, "Topología de arranque (ver hoja TOPOLOGIAS)."),
    ("LONGITUD_m", "TENDIDO", False, 12, None, "Longitud del tendido en metros."),
    ("CABLE_TIPO", "TENDIDO", False, 12, None, "Tipo de cable (texto libre, ej. tetra_sep)."),
    ("MATERIAL", "TENDIDO", False, 12, "Cobre", "Cobre o Aluminio (aluminio desde 16 mm²)."),
    ("AISLACION", "TENDIDO", False, 11, "PVC", "PVC (70 °C) o XLPE (90 °C)."),
    ("CLASE", "TENDIDO", False, 8, 5, "Clase 5 (flexible): Iz × 0,95. Clase 2 (rígido): sin reducción."),
    ("METODO", "INSTALACIÓN", False, 44, M[0], "Forma de instalar el cable; define la tabla 771.16."),
    ("N_TERNAS", "INSTALACIÓN", False, 10, 1, "Conductores en paralelo por fase."),
    ("AGRUP_CIRCUITOS", "INSTALACIÓN", False, 16, 1, "Circuitos agrupados en la misma canalización, contando este."),
    ("TEMP_AMBIENTE_C", "INSTALACIÓN", False, 16, 40, "Temperatura ambiente [°C]; base 40."),
    ("TEMP_SUELO_C", "INSTALACIÓN", False, 14, 25, "Temperatura del terreno [°C]; base 25 (solo enterrados)."),
    ("RESIST_TERMICA_KMW", "INSTALACIÓN", False, 19, 1.0, "Resistividad térmica del terreno [K·m/W]; base 1,0 (solo enterrados)."),
    ("IZ_A", "INSTALACIÓN", False, 9, None, "Opcional: Iz base por terna [A] del fabricante; reemplaza la tabla."),
    ("KIND", "PROTECCIÓN", False, 9, None, "Dispositivo: auto, mcb, mccb, rcbo o rccb."),
    ("NOTAS", "PROTECCIÓN", False, 28, None, "Observaciones."),
]
TABLERO_CAMPOS = [  # (campo, obligatorio, ayuda)
    ("NOMBRE", True, "Nombre del tablero"), ("TENSION_V", True, "Tensión de línea [V]"),
    ("SISTEMA", True, "Esquema de puesta a tierra (TN-S, TN-C-S, TT, IT)"), ("ICC_KA", True, "Icc en barras [kA]"),
    ("FP_DEFAULT", False, "Factor de potencia por defecto"), ("FRECUENCIA_HZ", False, "Frecuencia [Hz]"),
    ("PROYECTO", False, "Código del proyecto"), ("LONG_ACOMETIDA_m", False, "Largo de la acometida [m]"),
]


def _f(bold=False, color=TINTA, size=10, italic=False):
    return Font(name=FUENTE, bold=bold, color=color, size=size, italic=italic)


def _relleno(c):
    return PatternFill("solid", start_color=c, end_color=c)


FINA = Side(style="thin", color=LINEA)
BORDE = Border(bottom=FINA)
CAJA = Border(top=FINA, bottom=FINA, left=FINA, right=FINA)


def _titulo(ws, texto, sub=None):
    ws["A1"] = texto
    ws["A1"].font = _f(True, size=14)
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub
        ws["A2"].font = _f(color=MEDIO, size=9, italic=True)
    ws.sheet_view.showGridLines = False


def _hoja_tablero(wb, datos):
    ws = wb.create_sheet("TABLERO")
    _titulo(ws, "Datos del tablero", "Columna VALOR: lo que lee el motor. Los campos oscuros son obligatorios.")
    for i, h in enumerate(("CAMPO", "VALOR", "EJEMPLO", "DESCRIPCIÓN"), 1):
        c = ws.cell(4, i, h)
        c.font, c.fill = _f(True, "FFFFFF"), _relleno(OSCURO)
        c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[4].height = 22
    for r, (campo, obl, ayuda) in enumerate(TABLERO_CAMPOS, 5):
        d = datos.get(campo, (None, None))
        ws.cell(r, 1, campo).font = _f(True, OSCURO if obl else MEDIO)
        v = ws.cell(r, 2, d[0])
        v.font, v.fill, v.border = _f(True, ACENTO), _relleno(SUAVE), CAJA
        ws.cell(r, 3, d[1]).font = _f(color=MEDIO, italic=True)
        ws.cell(r, 4, ayuda + ("" if obl else " (opcional)")).font = _f(color=MEDIO)
        for c in range(1, 5):
            ws.cell(r, c).alignment = Alignment(vertical="center", indent=1, horizontal="left")
            ws.cell(r, c).border = BORDE if c != 2 else CAJA
        ws.row_dimensions[r].height = 20
    for col, w in zip("ABCD", (22, 36, 34, 52)):
        ws.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"TN-S,TN-C-S,TT,IT"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B7")


def _hoja_circuitos(wb, titulo, filas):
    ws = wb.create_sheet("CIRCUITOS")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Circuitos"
    ws["A1"].font = _f(True, size=14)
    ws.row_dimensions[1].height = 26
    HDR, R0 = 3, 4
    # bandas de grupo (fila 2)
    grupos = []
    for i, (n, g, *_r) in enumerate(COLUMNAS, 1):
        if not grupos or grupos[-1][0] != g:
            grupos.append([g, i, i])
        else:
            grupos[-1][2] = i
    for g, a, b in grupos:
        ws.merge_cells(start_row=2, start_column=a, end_row=2, end_column=b)
        c = ws.cell(2, a, g)
        c.font, c.alignment = _f(True, MEDIO, 8), Alignment(horizontal="left", vertical="center", indent=1)
        for k in range(a, b + 1):
            ws.cell(2, k).fill = _relleno(CLARO if grupos.index([g, a, b]) % 2 == 0 else "F3F4F6")
    ws.row_dimensions[2].height = 16
    for i, (n, g, obl, w, dflt, ayuda) in enumerate(COLUMNAS, 1):
        c = ws.cell(HDR, i, n)
        c.font, c.fill = _f(True, "FFFFFF", 9), _relleno(OSCURO if obl else MEDIO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.comment = Comment(f"{n}{' (obligatoria)' if obl else ''}\n{ayuda}", "AEA 90364")
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[HDR].height = 30
    idx = {n: i for i, (n, *_r) in enumerate(COLUMNAS, 1)}
    for k, fila in enumerate(filas):
        r = R0 + k
        for n, i in idx.items():
            v = fila.get(n)
            if v is None:
                v = next(c[4] for c in COLUMNAS if c[0] == n)
            c = ws.cell(r, i, v)
            c.font, c.border = _f(size=10, bold=(n == "TAG")), BORDE
            centrado = n not in ("DESCRIPCION", "DESTINO", "NOTAS", "METODO", "CABLE_TIPO")
            c.alignment = Alignment(horizontal="center" if centrado else "left", vertical="center", indent=0 if centrado else 1)
        ws.row_dimensions[r].height = 18
    ult = R0 + max(len(filas), 60)
    for n, vals in LISTAS.items():
        col = L(idx[n])
        if n == "METODO":
            f1 = f"=LISTAS!$B$2:$B${len(M) + 1}"
        else:
            f1 = '"' + ",".join(vals) + '"'
        dv = DataValidation(type="list", formula1=f1, allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{col}{R0}:{col}{ult}")
    ws.freeze_panes = ws.cell(R0, 3)
    ws.auto_filter.ref = f"A{HDR}:{L(len(COLUMNAS))}{R0 + len(filas) - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{HDR - 1}:{HDR}"


def _hoja_listas(wb):
    ws = wb.create_sheet("LISTAS")
    ws.sheet_view.showGridLines = False
    for i, h in enumerate(("CÓDIGO", "MÉTODO DE INSTALACIÓN", "TABLA 771.16"), 1):
        ws.cell(1, i, h)
    tablas = {"B1": "I", "B2": "III", "C": "III", "E": "III", "F": "III", "G": "III", "D1": "V", "D2": "VI"}
    for r, (cod, desc) in enumerate(A.METODOS_DESCRIPCION.items(), 2):
        ws.cell(r, 1, cod)
        ws.cell(r, 2, desc)
        ws.cell(r, 3, tablas[cod])
    for c in range(1, 4):
        h = ws.cell(1, c)
        h.font, h.fill = _f(True, "FFFFFF", 9), _relleno(OSCURO)
        h.alignment = Alignment(vertical="center", indent=1)
        for r in range(2, len(M) + 2):
            x = ws.cell(r, c)
            x.font, x.border = _f(), BORDE
            x.alignment = Alignment(vertical="center", indent=1)
    ws.cell(len(M) + 4, 1, "Cobre / Aluminio · PVC (70 °C) / XLPE (90 °C) · Clase 5 (flexible, Iz×0,95) / Clase 2 (rígido)").font = _f(color=MEDIO, size=9, italic=True)
    for col, w in zip("ABC", (10, 62, 14)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


def _hoja_leeme(wb):
    ws = wb.active
    ws.title = "LEEME"
    _titulo(ws, "Motor de cálculo AEA 90364 - entrada de circuitos",
            "Complete TABLERO y CIRCUITOS. El motor calcula In, sección, ΔU, Icu y genera el JSON, la MC y la SP del alimentador principal.")
    ws["A4"] = "Cómo se completa"
    ws["A4"].font = _f(True, size=11)
    pasos = ["1. TABLERO: tensión, esquema de tierra e Icc en barras (campos oscuros).",
             "2. CIRCUITOS: una fila por circuito; las columnas oscuras son obligatorias, las grises mejoran el cálculo.",
             "3. Elija material, aislación, clase y método de instalación de las listas; las celdas vacías usan el valor por defecto.",
             "4. PARENT_TAG cuelga un circuito de otro; si la carga del padre está vacía se suma la de sus hijos.",
             "5. Ejecute: python3 -B aea_pipeline.py input_motor_AEA_V2.xlsx"]
    for r, t in enumerate(pasos, 5):
        ws.cell(r, 1, t).font = _f()
    r = 5 + len(pasos) + 1
    ws.cell(r, 1, "Columnas de CIRCUITOS").font = _f(True, size=11)
    r += 1
    for i, h in enumerate(("COLUMNA", "GRUPO", "OBLIGATORIA", "DESCRIPCIÓN"), 1):
        c = ws.cell(r, i, h)
        c.font, c.fill = _f(True, "FFFFFF", 9), _relleno(OSCURO)
        c.alignment = Alignment(vertical="center", indent=1)
    for n, g, obl, w, dflt, ayuda in COLUMNAS:
        r += 1
        vals = (n, g.capitalize(), "Sí" if obl else "", ayuda + (f" Defecto: {dflt}." if dflt is not None else ""))
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = _f(i == 1, TINTA if i == 1 else MEDIO if i == 2 else TINTA), BORDE
            c.alignment = Alignment(vertical="center", indent=1)
    for col, w in zip("ABCD", (24, 18, 13, 96)):
        ws.column_dimensions[col].width = w


def _hoja_topologias(wb, origen):
    ws = wb.create_sheet("TOPOLOGIAS")
    _titulo(ws, "Topologías de arranque (campo TOPOLOGY)")
    filas = [r for r in origen.iter_rows(min_row=2, max_row=10, max_col=4, values_only=True)]
    for r, fila in enumerate(filas, 4):
        for c, v in enumerate(fila, 1):
            x = ws.cell(r, c, v)
            if r == 4:
                x.font, x.fill = _f(True, "FFFFFF", 9), _relleno(OSCURO)
            else:
                x.font = _f(r > 4 and c == 1 and True)
                x.border = BORDE
            x.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    for col, w in zip("ABCD", (14, 46, 48, 16)):
        ws.column_dimensions[col].width = w


def main(origen, destino):
    src = openpyxl.load_workbook(origen)
    hoja = src["CIRCUITOS"]
    hdr = next(r for r in range(1, 20) if hoja.cell(r, 1).value == "TAG")
    cols = [c.value for c in hoja[hdr]]
    filas = []
    for fila in hoja.iter_rows(min_row=hdr + 1, values_only=True):
        if fila and fila[0] not in (None, ""):
            filas.append(dict(zip(cols, fila)))
    tablero = {}
    for f in src["TABLERO"].iter_rows(min_row=1, max_col=3, values_only=True):
        if f[0] and f[1] is not None and str(f[0]).replace("_", "").isalnum() and f[0] != "CAMPO":
            tablero[f[0]] = (f[1], f[2])
    wb = openpyxl.Workbook()
    _hoja_leeme(wb)
    _hoja_tablero(wb, tablero)
    _hoja_circuitos(wb, "Circuitos", filas)
    _hoja_listas(wb)
    _hoja_topologias(wb, src["TOPOLOGIAS"])
    wb.save(destino)
    print("OK", destino, len(filas), "circuitos")


if __name__ == "__main__":
    main(*sys.argv[1:3])
